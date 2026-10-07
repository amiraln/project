"""Журналы и записи в Word (.docx) — тот же вид, что у печатных форм (print_list.html, print_detail.html)."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FONT = "Times New Roman"
EMPTY_ROWS = 10  # пустой журнал — с пустыми строками, чтобы заполнять от руки, как на печати


def _document(landscape):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Mm(12)
    section.left_margin = section.right_margin = Mm(10)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    return doc


def _paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=None, bold=False, upper=False, space_after=0,
               grey=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.all_caps = upper
    if size:
        run.font.size = Pt(size)
    if grey:
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return p


def _cell(cell, text, size=10, bold=False, center=False):
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)  # «\n» в тексте становится переносом строки
    run.bold = bold
    run.font.size = Pt(size)


def _repeat_on_each_page(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def _column_widths(doc, headers, rows, numbered):
    """Ширина граф по содержимому: «№ п/п» узкая, у остальных — по длине заголовка и записей, без крайностей,
    но не уже самого длинного слова графы, чтобы слова не разрывались."""
    section = doc.sections[0]
    number = Mm(12) if numbered else 0
    free = section.page_width - section.left_margin - section.right_margin - number
    weights, minimums = [], []
    for i in range(numbered, len(headers)):
        lines = [headers[i], *(line for cells in rows for line in cells[i].split("\n"))]
        weights.append(min(max(max(map(len, lines)), 8), 60) ** 0.5)
        minimums.append(Mm(2 * max(len(word) for line in lines for word in [*line.split(), ""]) + 3))
    widths = [free * w / sum(weights) for w in weights]
    narrow = [i for i, (w, low) in enumerate(zip(widths, minimums)) if w < low]
    if narrow and sum(minimums[i] for i in narrow) < free:  # узким — их минимум, остальное делят прочие графы
        rest = [i for i in range(len(widths)) if i not in narrow]
        left = free - sum(minimums[i] for i in narrow)
        for i in narrow:
            widths[i] = minimums[i]
        for i in rest:
            widths[i] = left * weights[i] / sum(weights[j] for j in rest)
    return ([number] if numbered else []) + [int(w) for w in widths]


def _table(doc, headers, rows, numbered):
    """Таблица журнала: заголовки, строка с номерами граф и записи; numbered — первая графа «№ п/п»."""
    if numbered:
        headers = ["№ п/п", *headers]
        rows = [[str(n), *cells] for n, cells in enumerate(rows, 1)]
    rows = rows or [[""] * len(headers) for _ in range(EMPTY_ROWS)]
    table = doc.add_table(rows=2 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = _column_widths(doc, headers, rows, numbered)
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width  # Word берёт ширину из каждой ячейки
    for i, header in enumerate(headers):
        _cell(table.rows[0].cells[i], header, size=9.5, bold=True, center=True)
        _cell(table.rows[1].cells[i], str(i + 1), size=8, center=True)
    _repeat_on_each_page(table.rows[0])
    _repeat_on_each_page(table.rows[1])
    for row, cells in zip(table.rows[2:], rows):
        for cell, text in zip(row.cells, cells):
            _cell(cell, text)
    return table


def _signatures(doc, org_head):
    p = _paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(24)
    p.add_run(f"Медицинская сестра ____________________\t\tРуководитель ____________________ {org_head}".rstrip())


def journal_document(sheet):
    """Журнал по данным печатной формы (views._journal_sheet)."""
    journal = sheet["journal"]
    doc = _document(journal.landscape)
    if journal.form:
        _paragraph(doc, journal.form, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _paragraph(doc, sheet["org_name"], size=11)
    _paragraph(doc, journal.title, size=14, bold=True, upper=True, space_after=4).paragraph_format.space_before = Pt(12)
    if sheet["group"]:
        _paragraph(doc, f"Группа: {sheet['group']}", bold=True)
    _paragraph(doc, f"Начат {sheet['started']}     Окончен {sheet['finished']}")
    if sheet["shown_from"] or sheet["shown_to"]:
        period = "".join([f" с {sheet['shown_from']:%d.%m.%Y}" if sheet["shown_from"] else "",
                          f" по {sheet['shown_to']:%d.%m.%Y}" if sheet["shown_to"] else ""])
        _paragraph(doc, f"Записи за период{period}")
    _paragraph(doc, "", space_after=6)
    _table(doc, sheet["headers"], [cells for _, cells in sheet["rows"]], numbered=not journal.form)
    if journal.footnote:
        _paragraph(doc, f"Примечание: {journal.footnote}", align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5).paragraph_format.space_before = Pt(6)
    if not journal.form:  # в утверждённых формах подписи — в самих графах
        _signatures(doc, sheet["org_head"])
    _paragraph(doc, f"Сформировано {sheet['printed_at']:%d.%m.%Y %H:%M}", align=WD_ALIGN_PARAGRAPH.LEFT, size=8,
               grey=True).paragraph_format.space_before = Pt(6)
    return doc


def record_document(sheet):
    """Одна запись целиком по данным печатной формы записи (views._record_sheet)."""
    doc = _document(landscape=False)
    _paragraph(doc, sheet["org_name"], size=11)
    _paragraph(doc, sheet["record_title"], size=14, bold=True, upper=True, space_after=8).paragraph_format.space_before = Pt(12)
    table = doc.add_table(rows=len(sheet["fields"]), cols=2)
    table.style = "Table Grid"
    for row, (name, value) in zip(table.rows, sheet["fields"]):
        _cell(row.cells[0], name, bold=True)
        _cell(row.cells[1], value)
        row.cells[0].width, row.cells[1].width = Mm(70), Mm(120)
    for inline in sheet["inlines"]:
        _paragraph(doc, inline["title"], size=12, bold=True, upper=True, space_after=4).paragraph_format.space_before = Pt(12)
        _table(doc, inline["headers"], inline["rows"], numbered=True)
    _signatures(doc, sheet["org_head"])
    _paragraph(doc, f"Сформировано {sheet['printed_at']:%d.%m.%Y %H:%M}", align=WD_ALIGN_PARAGRAPH.LEFT, size=8,
               grey=True).paragraph_format.space_before = Pt(6)
    return doc
