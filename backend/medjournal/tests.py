import datetime as dt
import io
from urllib.parse import quote

from docx import Document

from django.contrib.auth.models import Permission, User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from payments.models import Child, Guardian

from . import models as m
from . import views, word
from .journals import JOURNALS

D = "2026-09-20"
DT = "2026-09-20T09:30"
T = "10:15"

PAYLOADS = {
    "infectious": {"date_detected": D, "diagnosis": "Ветряная оспа"},
    "somatic": {"date_start": D, "date_end": "2026-09-24", "diagnosis": "ОРВИ"},
    "contacts": {"disease": "Ветряная оспа", "contact_date": D},
    "vaccinations": {"vaccine": "MMR", "dose": "V1", "date": D},
    "exemptions": {"kind": "MED", "vaccine": "DTP", "date_from": D, "reason": "ОРВИ"},
    "mantoux": {"test_type": "MANTOUX", "purpose": "SCREEN", "date_placed": D, "papule_mm": "12",
                "hyperemia_mm": "14", "previous_papule_mm": "5"},
    "tb-risk": {"risk_factor": "FREQ_ILL", "date_included": D},
    "phthisio": {"reason": "TURN", "referral_date": D, "admitted": "unknown"},
    "chemoprophylaxis": {"indication": "Вираж", "drug": "Изониазид", "dose": "10 мг/кг 1 р/д", "date_start": D},
    "helminths": {"exam_date": D, "method": "SCRAPE"},
    "health-passport": {"health_group": "II"},
    "risk-groups": {"category": "ALLERGY", "date_included": D},
    "food-brakerage": {"received_at": DT, "product": "Молоко 2,5%", "quantity": "20 л",
                       "organoleptic": "Соответствует", "responsible": "Иванова"},
    "kitchen-staff": {"date": D, "employee": "Сейтова А.", "position": "Повар",
                      "skin_ok": "on", "throat_ok": "on", "gut_ok": "on", "family_ok": "on",
                      "admitted": "on", "inspected_by": "Медсестра"},
    "food-norms": {"period_start": "2026-09-01", "period_end": "2026-09-10", "age_group": "GARDEN",
                   "child_days": "200", "responsible": "Медсестра"},
    "medical-cards": {"admission_date": D},
    "dish-quality": {"made_at": DT, "dish": "Суп гороховый", "organoleptic": "Без замечаний, готов",
                     "permitted_at": T, "executor": "Сейтова А., повар", "inspected_by": "Иванова"},
    "vitamin-c": {"prepared_at": DT, "dish": "Компот", "total_mg": "1250", "portion_mg": "50",
                  "responsible": "Медсестра"},
}


def inline_mgmt(journal, rows=None):
    data = {}
    for inl in journal.inlines:
        p = inl.model._meta.model_name
        data.update({f"{p}-TOTAL_FORMS": "0", f"{p}-INITIAL_FORMS": "0"})
    for p, items in (rows or {}).items():
        data[f"{p}-TOTAL_FORMS"] = str(len(items))
        for i, item in enumerate(items):
            data.update({f"{p}-{i}-{k}": v for k, v in item.items()})
    return data


class JournalsTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("nurse", "", "x", first_name="Айгерим", last_name="Касымова")
        self.client.force_login(self.user)
        self.child = Child.objects.create(p_id=1, fio="Нуров Алихан")
        Child.objects.create(p_id=2, fio="Выбывший", is_active=False)

    def test_all_journals_crud_and_print(self):
        self.assertEqual(set(PAYLOADS), {j.slug for j in JOURNALS})
        self.assertEqual(self.client.get(reverse("medjournal:index")).status_code, 200)
        for j in JOURNALS:
            with self.subTest(j.slug):
                data = dict(PAYLOADS[j.slug])
                if j.child_field:
                    data["child"] = str(self.child.pk)
                data.update(inline_mgmt(j))
                r = self.client.get(reverse("medjournal:add", args=[j.slug]))
                self.assertEqual(r.status_code, 200)
                r = self.client.post(reverse("medjournal:add", args=[j.slug]), data)
                if r.status_code != 302:
                    self.fail(f"{j.slug}: {r.context['form'].errors} "
                              f"{[fs.errors for _, fs in r.context['formsets']]}")
                obj = j.model.objects.get()
                for url in ("list", "print"):
                    r = self.client.get(reverse(f"medjournal:{url}", args=[j.slug]),
                                        {"date_from": "2026-09-01", "child": self.child.pk})
                    self.assertEqual(r.status_code, 200)
                    self.assertContains(r, "Нуров Алихан" if j.child_field else "") if j.child_field else None
                self.assertEqual(self.client.get(reverse("medjournal:edit", args=[j.slug, obj.pk])).status_code, 200)
                self.assertEqual(self.client.get(reverse("medjournal:record_print", args=[j.slug, obj.pk])).status_code, 200)
        # админка
        for j in JOURNALS:
            url = reverse(f"admin:medjournal_{j.model._meta.model_name}_changelist")
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_mantoux_result_and_turn(self):
        t = m.MantouxTest.objects.create(child=self.child, date_placed=dt.date(2026, 9, 1),
                                         papule_mm=12, previous_papule_mm=5)
        self.assertEqual(t.result, "POS")
        self.assertTrue(t.is_turn)
        t.papule_mm = 17; t.save()
        self.assertEqual(t.result, "HYPER")
        t.papule_mm = 0; t.hyperemia_mm = 3; t.save()
        self.assertEqual(t.result, "DOUBT")
        t.test_type = "DST"; t.hyperemia_mm = 0; t.save()
        self.assertEqual(t.result, "NEG")

    def test_inactive_children_hidden_and_signature_prefilled(self):
        r = self.client.get(reverse("medjournal:add", args=["somatic"]))
        self.assertNotContains(r, "Выбывший")
        r = self.client.get(reverse("medjournal:add", args=["vitamin-c"]))
        self.assertContains(r, "Айгерим Касымова")

    def test_new_record_is_filled_from_eorda_after_choosing_child(self):
        self.child.iin, self.child.address, self.child.contract_date = "210101500001", "ул. Абая, 5", dt.date(2025, 9, 1)
        self.child.save()
        Guardian.objects.create(child_pid=480001, child_iin="210101500001", gender=Guardian.MOTHER,
                                fio="Нурова Айгуль", phone="+77010000001, +77020000002")
        Guardian.objects.create(child_pid=1, gender=Guardian.FATHER, fio="Нуров Ерлан", phone="+77030000003")
        Guardian.objects.create(child_pid=1, fio="Пол неизвестен")
        expected = {"mother": "Нурова Айгуль, +77010000001, +77020000002", "father": "Нуров Ерлан, +77030000003",
                    "address": "ул. Абая, 5", "admission_date": "2025-09-01"}

        r = self.client.get(reverse("medjournal:add", args=["medical-cards"]))
        self.assertEqual(r.context["autofill"]["facts"][self.child.pk], expected)
        self.assertContains(r, 'id="autofill-data"')

        # из журнала, отфильтрованного по ребёнку, поля заполнены сразу
        r = self.client.get(reverse("medjournal:add", args=["medical-cards"]), {"child": self.child.pk})
        form = r.context["form"]
        self.assertEqual({name: form[name].value() for name in expected}, expected)
        self.assertContains(r, 'name="admission_date" value="2025-09-01"')

        # в журналах без этих полей и при изменении записи ничего не подставляется
        self.assertIsNone(self.client.get(reverse("medjournal:add", args=["somatic"])).context["autofill"])
        card = m.MedicalCard.objects.create(child=self.child)
        self.assertIsNone(self.client.get(reverse("medjournal:edit", args=["medical-cards", card.pk])).context["autofill"])

    def test_fill_all_creates_records_with_known_fields_only(self):
        self.assertEqual([j.slug for j in JOURNALS if views._fillable(j)], ["health-passport", "medical-cards"])
        self.child.address, self.child.contract_date = "ул. Абая, 5", dt.date(2025, 9, 1)
        self.child.save()
        Guardian.objects.create(child_pid=1, gender=Guardian.MOTHER, fio="Нурова Айгуль", phone="+77010000001")
        second = Child.objects.create(p_id=3, fio="Серикова Дана")  # без данных из e-orda
        third = Child.objects.create(p_id=4, fio="Ахметов Тимур", address="пр. Мира, 10")
        m.MedicalCard.objects.create(child=third, mother="Вписано вручную")
        url = reverse("medjournal:fill_all", args=["medical-cards"])

        self.assertContains(self.client.get(reverse("medjournal:list", args=["medical-cards"])), url)
        self.assertNotContains(self.client.get(reverse("medjournal:list", args=["somatic"])), "Заполнить всех")
        r = self.client.get(url)
        self.assertEqual((r.context["new"], r.context["changed"]), (2, 1))
        self.assertFalse(m.MedicalCard.objects.filter(child=self.child).exists())  # пока только подтверждение

        self.assertRedirects(self.client.post(url), reverse("medjournal:list", args=["medical-cards"]))
        cards = {c.child_id: c for c in m.MedicalCard.objects.all()}
        self.assertEqual(set(cards), {self.child.pk, second.pk, third.pk})  # выбывшего нет
        card = cards[self.child.pk]
        self.assertEqual((card.mother, card.father, card.address, card.admission_date, card.polyclinic, card.created_by),
                         ("Нурова Айгуль, +77010000001", "", "ул. Абая, 5", dt.date(2025, 9, 1), "", self.user))
        self.assertEqual((cards[second.pk].address, cards[second.pk].admission_date), ("", None))
        self.assertEqual((cards[third.pk].mother, cards[third.pk].address), ("Вписано вручную", "пр. Мира, 10"))

        r = self.client.get(url)
        self.assertEqual((r.context["new"], r.context["changed"]), (0, 0))
        self.client.post(url)
        self.assertEqual(m.MedicalCard.objects.count(), 3)

        # паспорт здоровья: заводится на каждого ребёнка, остальные поля пустые
        self.client.post(reverse("medjournal:fill_all", args=["health-passport"]))
        self.assertEqual(m.HealthPassport.objects.filter(health_group="").count(), 3)

    def test_fill_all_is_not_offered_where_records_need_manual_data(self):
        self.assertEqual(self.client.get(reverse("medjournal:fill_all", args=["infectious"])).status_code, 404)
        self.assertEqual(self.client.post(reverse("medjournal:fill_all", args=["vitamin-c"])).status_code, 404)
        u = User.objects.create_user("viewer", "", "x", is_staff=True)
        u.user_permissions.add(*Permission.objects.filter(codename__in=["view_medicalcard", "add_medicalcard"]))
        self.client.force_login(u)
        self.assertEqual(self.client.post(reverse("medjournal:fill_all", args=["medical-cards"])).status_code, 403)
        self.assertFalse(m.MedicalCard.objects.exists())

    def test_print_has_organization_and_journal_dates(self):
        url = reverse("medjournal:print", args=["somatic"])
        r = self.client.get(url)
        self.assertContains(r, "Детский сад «Ботаканым»")
        self.assertNotContains(r, "Ясли")
        self.assertContains(r, "Начат <span id=\"started\">«___» __________ 20__ г.</span>")
        self.assertContains(r, 'name="started"')

        save = reverse("medjournal:print_period", args=["somatic"])
        r = self.client.post(save, {"started": "2025-09-01", "finished": ""})
        self.assertEqual(r.json(), {"ok": True, "started": "«01» сентября 2025 г.", "finished": "«___» __________ 20__ г."})
        r = self.client.get(url, {"date_from": "2026-09-01"})
        self.assertContains(r, "«01» сентября 2025 г.")
        self.assertContains(r, "Записи за период с 01.09.2026")
        self.assertContains(r, 'value="2025-09-01"')
        self.assertNotContains(self.client.get(reverse("medjournal:print", args=["infectious"])), "сентября 2025")

        r = self.client.post(save, {"started": "2025-09-01", "finished": "2025-08-31"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(m.JournalPeriod.objects.get().finished, None)
        self.assertContains(self.client.get(reverse("medjournal:record_print", args=["medical-cards", m.MedicalCard.objects.create(child=self.child).pk])),
                            "Детский сад «Ботаканым»")

        # без права изменять журнал даты видны, но поменять их нельзя
        u = User.objects.create_user("viewer", "", "x", is_staff=True)
        u.user_permissions.add(Permission.objects.get(codename="view_somaticdiseaserecord"))
        self.client.force_login(u)
        r = self.client.get(url)
        self.assertContains(r, "«01» сентября 2025 г.")
        self.assertNotContains(r, 'name="started"')
        self.assertEqual(self.client.post(save, {"started": "2020-01-01"}).status_code, 403)

    def test_period_creates_a_record_for_each_day(self):
        url = reverse("medjournal:add", args=["kitchen-staff"])
        staff = {**PAYLOADS["kitchen-staff"], "date": "2026-09-18"}  # пятница
        r = self.client.post(url, {**staff, "period-until": "2026-09-22", "period-skip_weekends": "on"})
        self.assertRedirects(r, reverse("medjournal:list", args=["kitchen-staff"]))
        days = sorted(m.KitchenStaffInspection.objects.values_list("date", flat=True))
        self.assertEqual(days, [dt.date(2026, 9, 18), dt.date(2026, 9, 21), dt.date(2026, 9, 22)])
        self.assertEqual(set(m.KitchenStaffInspection.objects.values_list("employee", "created_by")),
                         {("Сейтова А.", self.user.pk)})

        m.KitchenStaffInspection.objects.all().delete()
        self.client.post(url, {**staff, "period-until": "2026-09-22"})  # с выходными
        self.assertEqual(m.KitchenStaffInspection.objects.count(), 5)

        # дата со временем: время у всех записей то же
        self.client.post(reverse("medjournal:add", args=["food-brakerage"]),
                         {**PAYLOADS["food-brakerage"], "period-until": "2026-09-22"})
        times = [timezone.localtime(t) for t in m.PerishableFoodBrakerage.objects.order_by("received_at")
                 .values_list("received_at", flat=True)]
        self.assertEqual([(t.day, t.hour, t.minute) for t in times], [(20, 9, 30), (21, 9, 30), (22, 9, 30)])

        # у записей с ребёнком — тот же ребёнок, результат пробы считается в каждой
        self.client.post(reverse("medjournal:add", args=["mantoux"]),
                         {**PAYLOADS["mantoux"], "child": self.child.pk, "period-until": "2026-09-21"})
        self.assertEqual(list(m.MantouxTest.objects.values_list("child", "result")), [(self.child.pk, "POS")] * 2)

    def test_period_errors_save_nothing(self):
        url = reverse("medjournal:add", args=["vitamin-c"])
        for extra, field in (({"period-until": "2026-09-19"}, "until"),  # раньше даты записи (20.09)
                             ({"period-until": "2027-12-31"}, "until"),  # больше года
                             ({"period-until": "2026-09-20", "period-skip_weekends": "on"}, "skip_weekends")):  # только вс
            r = self.client.post(url, {**PAYLOADS["vitamin-c"], **extra})
            self.assertEqual(r.status_code, 200, extra)
            self.assertIn(field, r.context["period"].errors, extra)
        self.assertFalse(m.VitaminCRecord.objects.exists())

        # период — только в новых записях журналов, где дата обязательна и нет вложенных таблиц
        self.assertIsNotNone(self.client.get(url).context["period"])
        for slug in ("chemoprophylaxis", "food-norms", "medical-cards", "health-passport"):
            self.assertIsNone(self.client.get(reverse("medjournal:add", args=[slug])).context["period"], slug)
        record = m.VitaminCRecord.objects.create(prepared_at=timezone.now(), dish="Компот", total_mg=1250,
                                                 portion_mg=50, responsible="Медсестра")
        self.assertIsNone(self.client.get(reverse("medjournal:edit", args=["vitamin-c", record.pk])).context["period"])

    def somatic(self, child, diagnosis="ОРВИ", day=20):
        return m.SomaticDiseaseRecord.objects.create(child=child, date_start=dt.date(2026, 9, day), diagnosis=diagnosis)

    def test_child_journals_are_split_by_groups(self):
        self.child.group_name = "Балапан"
        self.child.save()
        other = Child.objects.create(p_id=3, fio="Серикова Дана", group_name="Ромашка")
        self.somatic(self.child, "Бронхит")
        self.somatic(other, "Отит")
        url = reverse("medjournal:list", args=["somatic"])

        r = self.client.get(url)
        self.assertEqual([(t["label"], t["count"]) for t in r.context["group_tabs"]],
                         [("Все группы", 2), ("Балапан", 1), ("Ромашка", 1)])
        self.assertEqual(r.context["headers"][:2], ["Ребёнок", "Группа"])  # во всех группах — колонка группы

        r = self.client.get(url, {"group": "Ромашка"})
        self.assertEqual([obj.diagnosis for obj, _ in r.context["rows"]], ["Отит"])
        self.assertNotIn("Группа", r.context["headers"])
        self.assertEqual([g for g, _ in r.context["children"]], ["Ромашка"])  # в фильтре — дети этой группы
        self.assertContains(r, 'name="group" value="Ромашка"')

        r = self.client.get(reverse("medjournal:print", args=["somatic"]), {"group": "Балапан"})
        self.assertContains(r, "Группа: Балапан")
        self.assertContains(r, "Бронхит")
        self.assertNotContains(r, "Отит")

        # в форме дети сгруппированы; в журналах без детей групп нет
        self.assertContains(self.client.get(reverse("medjournal:add", args=["somatic"])), '<optgroup label="Ромашка">')
        self.assertIsNone(self.client.get(reverse("medjournal:list", args=["vitamin-c"])).context["group_tabs"])

    def test_bulk_delete_selected_or_all_by_filter(self):
        self.child.group_name = "Балапан"
        self.child.save()
        other = Child.objects.create(p_id=3, fio="Серикова Дана", group_name="Ромашка")
        a, b, c = self.somatic(self.child), self.somatic(self.child), self.somatic(other)
        url = reverse("medjournal:bulk", args=["somatic"])

        r = self.client.post(url, {"action": "delete", "ids": [a.pk, b.pk], "query": "group=Балапан"})
        self.assertEqual(r.context["count"], 2)
        self.assertEqual(m.SomaticDiseaseRecord.objects.count(), 3)  # пока только подтверждение
        r = self.client.post(url, {"action": "delete", "ids": r.context["ids"], "query": "group=Балапан", "confirm": "1"})
        self.assertRedirects(r, reverse("medjournal:list", args=["somatic"]) + "?group=Балапан",
                             fetch_redirect_response=False)
        self.assertEqual(list(m.SomaticDiseaseRecord.objects.values_list("pk", flat=True)), [c.pk])

        d = self.somatic(self.child)
        r = self.client.post(url, {"action": "delete", "all": "1", "query": "group=Ромашка"})
        self.assertEqual(r.context["ids"], str(c.pk))  # «все по фильтру» — только записи группы
        self.client.post(url, {"action": "delete", "ids": r.context["ids"], "confirm": "1"})
        self.assertEqual(list(m.SomaticDiseaseRecord.objects.values_list("pk", flat=True)), [d.pk])

        u = User.objects.create_user("editor", "", "x", is_staff=True)
        u.user_permissions.add(*Permission.objects.filter(codename__in=["view_somaticdiseaserecord",
                                                                         "change_somaticdiseaserecord"]))
        self.client.force_login(u)
        r = self.client.post(url, {"action": "delete", "ids": d.pk, "confirm": "1"})
        self.assertEqual(r.status_code, 403)
        self.assertTrue(m.SomaticDiseaseRecord.objects.exists())

    def test_bulk_edit_changes_one_field_in_all_selected(self):
        a, b, c = self.somatic(self.child), self.somatic(self.child), self.somatic(self.child, "Отит")
        url = reverse("medjournal:bulk", args=["somatic"])
        r = self.client.post(url, {"action": "edit", "ids": f"{a.pk},{b.pk}"})
        self.assertIn("diagnosis", r.context["form"].fields)
        self.assertNotIn("child", r.context["form"].fields)

        body = {"action": "edit", "ids": f"{a.pk},{b.pk}", "confirm": "1", "field": "date_end", "date_end": "2026-09-25"}
        self.assertEqual(self.client.post(url, body).status_code, 302)
        ends = dict(m.SomaticDiseaseRecord.objects.values_list("pk", "date_end"))
        self.assertEqual(ends, {a.pk: dt.date(2026, 9, 25), b.pk: dt.date(2026, 9, 25), c.pk: None})

        # обязательное поле нельзя очистить; правило журнала (выздоровление не раньше начала) — ни одной записи
        r = self.client.post(url, {**body, "field": "diagnosis", "diagnosis": ""})
        self.assertTrue(r.context["value_errors"])
        r = self.client.post(url, {**body, "date_end": "2026-09-01"})
        self.assertEqual(r.context["record_errors_count"], 2)
        self.assertEqual(m.SomaticDiseaseRecord.objects.filter(date_end=dt.date(2026, 9, 25)).count(), 2)

        # запись сохраняется целиком: результат пробы Манту пересчитывается
        test = m.MantouxTest.objects.create(child=self.child, date_placed=dt.date(2026, 9, 1), papule_mm=3)
        self.client.post(reverse("medjournal:bulk", args=["mantoux"]),
                         {"action": "edit", "ids": test.pk, "confirm": "1", "field": "papule_mm", "papule_mm": "18"})
        test.refresh_from_db()
        self.assertEqual((test.papule_mm, test.result), (18, "HYPER"))

    def test_food_journals_print_as_official_forms(self):
        for slug in ("food-brakerage", "vitamin-c", "dish-quality"):
            data = {**PAYLOADS[slug], **inline_mgmt(next(j for j in JOURNALS if j.slug == slug))}
            self.assertEqual(self.client.post(reverse("medjournal:add", args=[slug]), data).status_code, 302, slug)

        def printed(slug):
            return self.client.get(reverse("medjournal:print", args=[slug])).content.decode()

        page = printed("food-brakerage")
        headers = ["Дата и час поступления продовольственного сырья и пищевых продуктов",
                   "Наименование пищевых продуктов",
                   "Количество поступившего продовольственного сырья и пищевых продуктов (в килограммах, литрах, штуках)",
                   "Результаты органолептической оценки поступившего продовольственного сырья и пищевых продуктов",
                   "Конечный срок реализации продовольственного сырья и пищевых продуктов",
                   "Дата и час фактической реализации продовольственного сырья и пищевых продуктов по дням",
                   "Ф.И.О. (при наличии) подпись ответственного лица", "Примечание *"]
        self.assertIn("".join(f"<th>{h}</th>" for h in headers), page)
        self.assertIn("".join(f"<th>{n}</th>" for n in range(1, 9)) + "</tr>", page)  # графы 1–8, без «№ п/п»
        self.assertIn("Форма 1", page)
        self.assertIn("Примечание: * Указываются факты списания, возврата продуктов", page)
        self.assertNotIn("№ п/п", page)
        self.assertNotIn("Медицинская сестра ____", page)

        page = printed("vitamin-c")
        headers = ["Дата и час приготовления блюда", "Наименование блюда", "Общее количество добавленного витамина",
                   "Содержание витамина «С» в одной порции", "Подпись ответственного лица"]
        self.assertIn("".join(f"<th>{h}</th>" for h in headers), page)
        for text in ("Форма 2", "Журнал «С-витаминизации»", "20.09.2026 09:30", "1250 мг", "50 мг"):
            self.assertIn(text, page)

        page = printed("dish-quality")
        headers = ["Дата, время изготовления блюд и кулинарных изделий", "Наименование блюд и кулинарных изделий",
                   "Органолептическая оценка, включая оценку степени готовности блюд и кулинарных изделий",
                   "Разрешение к реализации (время)",
                   "Ответственный исполнитель (Ф.И.О. (при его наличии), должность)",
                   "Ф.И.О. (при его наличии), лица проводившего бракераж", "Примечание"]
        self.assertIn("".join(f"<th>{h}</th>" for h in headers), page)
        for text in ("Форма 3", "10:15", "Сейтова А., повар",
                     "Примечание: в графе 7 указываются наименования готовой продукции, не допущенных к реализации"):
            self.assertIn(text, page)

        # остальные журналы печатаются как раньше
        page = printed("somatic")
        self.assertIn("№ п/п", page)
        self.assertNotIn("Форма ", page)

    def word(self, url, params=None):
        r = self.client.get(url, params or {})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], word.CONTENT_TYPE)
        return r, Document(io.BytesIO(b"".join(r.streaming_content)))

    def test_journal_downloads_as_word_like_its_printout(self):
        self.child.group_name = "Балапан"
        self.child.save()
        other = Child.objects.create(p_id=3, fio="Серикова Дана", group_name="Ромашка")
        m.SomaticDiseaseRecord.objects.create(child=self.child, date_start=dt.date(2026, 9, 20), diagnosis="Бронхит")
        m.SomaticDiseaseRecord.objects.create(child=other, date_start=dt.date(2026, 9, 21), diagnosis="Отит")
        m.JournalPeriod.objects.create(journal="somatic", started=dt.date(2025, 9, 1))

        r, doc = self.word(reverse("medjournal:word", args=["somatic"]), {"group": "Балапан"})
        self.assertIn("attachment", r["Content-Disposition"])
        self.assertIn(quote("2. Журнал соматической заболеваемости — Балапан"), r["Content-Disposition"])
        texts = [p.text for p in doc.paragraphs]
        for text in ("Детский сад «Ботаканым»", "Журнал соматической заболеваемости", "Группа: Балапан",
                     "Начат «01» сентября 2025 г.     Окончен «___» __________ 20__ г."):
            self.assertIn(text, texts)
        self.assertTrue(any(t.startswith("Медицинская сестра") for t in texts))
        table = doc.tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells][:3], ["№ п/п", "Ребёнок", "Дата начала заболевания"])
        self.assertEqual([c.text for c in table.rows[1].cells][:3], ["1", "2", "3"])
        self.assertEqual(len(table.rows), 3)  # заголовок, номера граф и одна запись группы
        self.assertEqual([c.text for c in table.rows[2].cells][:4], ["1", "Нуров Алихан", "20.09.2026", "Бронхит"])

        # утверждённая форма: «Форма 1», графы с 1 без «№ п/п», примечание к форме, без общей строки подписей
        _, doc = self.word(reverse("medjournal:word", args=["food-brakerage"]))
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(texts[0], "Форма 1")
        self.assertIn("Примечание: * Указываются факты списания, возврата продуктов", texts)
        self.assertFalse(any(t.startswith("Медицинская сестра") for t in texts))
        table = doc.tables[0]
        self.assertEqual(table.rows[0].cells[7].text, "Примечание *")
        self.assertEqual([c.text for c in table.rows[1].cells], [str(n) for n in range(1, 9)])
        self.assertEqual(len(table.rows), 2 + word.EMPTY_ROWS)  # пустой журнал — пустые строки для записи от руки
        self.assertEqual(doc.sections[0].orientation, 1)  # альбомная

        u = User.objects.create_user("teacher", "", "x", is_staff=True)
        self.client.force_login(u)
        self.assertEqual(self.client.get(reverse("medjournal:word", args=["somatic"])).status_code, 403)

    def test_record_downloads_as_word(self):
        card = m.MedicalCard.objects.create(child=self.child, mother="Нурова Айгуль, +77010000001",
                                            admission_date=dt.date(2025, 9, 1))
        m.MedicalCardEntry.objects.create(card=card, date=dt.date(2026, 9, 20), specialist="Педиатр",
                                          findings="Жалоб нет\nЗдоров")
        _, doc = self.word(reverse("medjournal:record_word", args=["medical-cards", card.pk]))
        fields = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
        self.assertEqual(fields["Ребёнок"], "Нуров Алихан")
        self.assertEqual(fields["Мать (ФИО, телефон)"], "Нурова Айгуль, +77010000001")
        self.assertEqual(fields["Дата поступления в д/с"], "01.09.2025")
        entries = doc.tables[1]
        self.assertEqual([c.text for c in entries.rows[0].cells][:3], ["№ п/п", "Дата", "Специалист / вид осмотра"])
        self.assertEqual(entries.rows[2].cells[3].text, "Жалоб нет\nЗдоров")  # перенос строки сохраняется
        self.assertIn("Записи осмотров", [p.text for p in doc.paragraphs])
        page = self.client.get(reverse("medjournal:record_print", args=["medical-cards", card.pk]))
        self.assertContains(page, reverse("medjournal:record_word", args=["medical-cards", card.pk]))

    def test_food_norms_inline_math(self):
        j = next(j for j in JOURNALS if j.slug == "food-norms")
        data = dict(PAYLOADS["food-norms"])
        data.update(inline_mgmt(j, {"foodnormline": [
            {"product": "Молоко", "norm_g": "350", "issued_total_g": "63000"},
            {"product": "Мясо", "norm_g": "60", "issued_total_g": "13200"},
        ]}))
        r = self.client.post(reverse("medjournal:add", args=["food-norms"]), data)
        self.assertEqual(r.status_code, 302)
        sheet = m.FoodNormSheet.objects.get()
        lines = list(sheet.lines.all())
        self.assertEqual(lines[0].actual_g, 315)
        self.assertEqual(lines[0].percent, 90)
        self.assertEqual(lines[1].percent, 110)
        r = self.client.get(reverse("medjournal:record_print", args=["food-norms", sheet.pk]))
        self.assertContains(r, "315")

    def test_validation_rules(self):
        staff = {**PAYLOADS["kitchen-staff"]}
        staff.pop("skin_ok")
        r = self.client.post(reverse("medjournal:add", args=["kitchen-staff"]), staff)
        self.assertIn("admitted", r.context["form"].errors)

    def test_permissions(self):
        u = User.objects.create_user("teacher", "", "x", is_staff=True)
        self.client.force_login(u)
        self.assertEqual(self.client.get(reverse("medjournal:list", args=["infectious"])).status_code, 403)
